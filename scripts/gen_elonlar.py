#!/usr/bin/env python3
"""Ilm-fan va innovatsiyalar — grant loyihalari (docs/*.xlsx) -> story-e'lonlar seed + JS fallback.

Ishga tushirish (uttp-next ichidan):  python3 scripts/gen_elonlar.py   (openpyxl kerak)

Manbalar:
  docs/oxirgi_3_yillikda_grant_loyihalar.xlsx   - TDTU grant loyihalari (8 ta, 2024–2026)
  docs/biofarmasevtika_grant_loyihalar.xlsx     - Milliy biofarmatsevtika instituti (7 ta)

Chiqish:
  uttp-next/supabase/seed_elonlar.sql   - public.elonlar qatorlari (Instagram-story ko'rinishidagi e'lonlar)
  uttp-next/lib/elonlar-fallback.js     - INITIAL_ELONLAR (Supabase bo'lmaganda)
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SRC_TDTU = ROOT / 'docs/oxirgi_3_yillikda_grant_loyihalar.xlsx'
SRC_BIO = ROOT / 'docs/biofarmasevtika_grant_loyihalar.xlsx'
OUT_SQL = ROOT / 'uttp-next/supabase/seed_elonlar.sql'
OUT_JS = ROOT / 'uttp-next/lib/elonlar-fallback.js'

ADMIN_ID = '44444444-4444-4444-4444-444444444444'   # seed.sql dagi super admin
RANGLAR = ['primary', 'teal', 'violet', 'navy', 'success', 'accent']


def norm(s) -> str:
    s = re.sub(r'\s+', ' ', str(s or '')).strip()
    return s.strip('“”"\'« »')


def qisqart(s: str, n: int) -> str:
    s = norm(s)
    if len(s) <= n:
        return s
    kes = s[: n - 1]
    if ' ' in kes:
        kes = kes[: kes.rfind(' ')]
    return kes + '…'


def q(v) -> str:
    if v is None:
        return 'NULL'
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def bosqich_va_muddat(matn: str, standart_muddat: str):
    m = norm(matn)
    yillar = re.findall(r'(20\d\d)\s*[-–]\s*(20\d\d)', m)
    muddat = f'{yillar[0][0]}–{yillar[0][1]}' if yillar else standart_muddat
    past = m.lower()
    if 'yakunlan' in past:
        bosqich = 'yakunlangan'
    elif 'davom' in past:
        bosqich = 'davom etmoqda'
    else:
        bosqich = 'rejalashtirilgan'
    return bosqich, muddat


elonlar = []

# ---- TDTU: satrlar 2..: nomi, mualliflar, soha, mexanizm, qiymat, bosqich ----
ws = openpyxl.load_workbook(SRC_TDTU, data_only=True).worksheets[0]
for row in list(ws.iter_rows(values_only=True))[2:]:
    nomi, mualliflar, soha, mexanizm, qiymat, bosqich = (list(row) + [None] * 6)[:6]
    if not nomi and not mualliflar:
        continue
    sarlavha = qisqart(nomi, 140) if nomi else f'Grant loyihasi — {norm(mualliflar)}'
    b, muddat = bosqich_va_muddat(bosqich, '2024–2026')
    elonlar.append({
        'sarlavha': sarlavha,
        'matn': qisqart(mexanizm or nomi or '', 420),
        'muammo': norm(soha) and f'Soha: {norm(soha)}. ' + (norm(nomi) if nomi else ''),
        'mexanizm': norm(mexanizm) or None,
        'soha': norm(soha) or None,
        'turi': 'grant',
        'muassasa': 'Toshkent davlat tibbiyot universiteti',
        'mualliflar': norm(mualliflar) or None,
        'qiymat': int(float(str(qiymat).replace(' ', ''))) if qiymat else None,
        'bosqich': b,
        'muddat': muddat,
        'manba': 'TDTU — oxirgi 3 yillik grant loyihalari',
    })

# ---- Biofarm: satrlar 5..: tr, nomi, mualliflar, muammo, mexanizm, qiymat, bosqich ----
ws = openpyxl.load_workbook(SRC_BIO, data_only=True).worksheets[0]
for row in list(ws.iter_rows(values_only=True))[5:]:
    _tr, nomi, mualliflar, muammo, mexanizm, qiymat, bosqich = (list(row) + [None] * 7)[:7]
    if not nomi:
        continue
    b, muddat = bosqich_va_muddat(bosqich, '2024–2026')
    elonlar.append({
        'sarlavha': qisqart(nomi, 140),
        'matn': qisqart(muammo or mexanizm or '', 420),
        'muammo': norm(muammo) or None,
        'mexanizm': norm(mexanizm) or None,
        'soha': 'farmatsevtika',
        'turi': 'grant',
        'muassasa': 'Milliy biofarmatsevtika ilmiy-tadqiqot instituti',
        'mualliflar': norm(mualliflar) or None,
        'qiymat': int(float(str(qiymat).replace(' ', ''))) if qiymat else None,
        'bosqich': b,
        'muddat': muddat,
        'manba': 'Milliy biofarmatsevtika instituti grant loyihalari',
    })

# ---- Platforma e'lonlari (namuna) ----
elonlar.append({
    'sarlavha': 'Ilm-fan va innovatsiyalar bo‘limi ishga tushdi',
    'matn': 'TDTU va Milliy biofarmatsevtika instituti grant loyihalari endi platformada. Har bir loyiha haqida qisqa e‘lon — story ko‘rinishida. Vazirlik va administrator yangi e‘lon joylashi mumkin.',
    'muammo': None, 'mexanizm': None, 'soha': None,
    'turi': 'elon', 'muassasa': 'ETTP', 'mualliflar': None, 'qiymat': None,
    'bosqich': None, 'muddat': None, 'manba': 'ETTP boshqaruvi',
})
elonlar.append({
    'sarlavha': 'Doktorantura 2026: 162 nafar qabul qilindi',
    'matn': 'TDTU doktoranturasiga 2026-yilda 162 nafar tayanch doktorant va DSc izlanuvchi qabul qilindi. Eng ko‘p — Stomatologiya, Nevrologiya va Ichki kasalliklar ixtisosliklari.',
    'muammo': None, 'mexanizm': None, 'soha': 'tibbiyot',
    'turi': 'innovatsiya', 'muassasa': 'Toshkent davlat tibbiyot universiteti', 'mualliflar': None, 'qiymat': None,
    'bosqich': None, 'muddat': '2026', 'manba': 'TDTU doktorantlar ro‘yxati, 04.09.2026',
})

for i, e in enumerate(elonlar):
    e['rang'] = RANGLAR[i % len(RANGLAR)]
    e['tartib'] = i + 1

# ---- SQL ----
COLS = ('sarlavha', 'matn', 'muammo', 'mexanizm', 'soha', 'turi', 'rang', 'muassasa', 'mualliflar', 'qiymat', 'bosqich', 'muddat', 'manba', 'muallif_id', 'holat', 'tartib')
lines = [
    '-- =========================================================',
    "-- ETTP — Ilm-fan va innovatsiyalar: story-e'lonlar (grant loyihalari)",
    f'-- Manba: docs/{SRC_TDTU.name}, docs/{SRC_BIO.name}',
    f"-- {len(elonlar)} e'lon. GENERATSIYA QILINGAN FAYL — scripts/gen_elonlar.py; qo‘lda tahrirlanmaydi.",
    '-- Idempotent: bir xil sarlavha bo‘lsa qayta kiritilmaydi.',
    '-- =========================================================',
    '',
    f'INSERT INTO public.elonlar ({", ".join(COLS)})',
    'SELECT v.* FROM (VALUES',
    ',\n'.join(
        '  (' + ', '.join(q(x) for x in (
            e['sarlavha'], e['matn'], e.get('muammo'), e.get('mexanizm'), e.get('soha'), e['turi'], e['rang'], e['muassasa'], e['mualliflar'], e['qiymat'],
            e['bosqich'], e['muddat'], e['manba'], ADMIN_ID, 'faol', e['tartib'])) + ')'
        for e in elonlar
    ),
    f') AS v({", ".join(COLS)})',
    "WHERE NOT EXISTS (SELECT 1 FROM public.elonlar e WHERE e.sarlavha = v.sarlavha);",
    '',
    '-- Ustun turlarini moslash (VALUES matn sifatida keladi)',
]
# VALUES'dagi muallif_id matn — uuid'ga cast kerak; qiymat integer. SELECT v.* o'rniga aniq castlar:
lines[7:] = [
    f'INSERT INTO public.elonlar ({", ".join(COLS)})',
    'SELECT v.sarlavha, v.matn, v.muammo, v.mexanizm, v.soha, v.turi, v.rang, v.muassasa, v.mualliflar, v.qiymat::numeric, v.bosqich, v.muddat, v.manba, v.muallif_id::uuid, v.holat, v.tartib::int',
    'FROM (VALUES',
    ',\n'.join(
        '  (' + ', '.join(q(x) for x in (
            e['sarlavha'], e['matn'], e.get('muammo'), e.get('mexanizm'), e.get('soha'), e['turi'], e['rang'], e['muassasa'], e['mualliflar'], e['qiymat'],
            e['bosqich'], e['muddat'], e['manba'], ADMIN_ID, 'faol', e['tartib'])) + ')'
        for e in elonlar
    ),
    f') AS v({", ".join(COLS)})',
    "WHERE NOT EXISTS (SELECT 1 FROM public.elonlar e WHERE e.sarlavha = v.sarlavha);",
    '',
]
OUT_SQL.write_text('\n'.join(lines), encoding='utf-8')

# ---- JS fallback ----
js_rows = []
for i, e in enumerate(elonlar):
    js_rows.append({
        'id': i + 1, 'sarlavha': e['sarlavha'], 'matn': e['matn'], 'muammo': e.get('muammo'), 'mexanizm': e.get('mexanizm'), 'soha': e.get('soha'), 'turi': e['turi'], 'rang': e['rang'],
        'muassasa': e['muassasa'], 'mualliflar': e['mualliflar'], 'qiymat': e['qiymat'], 'bosqich': e['bosqich'],
        'muddat': e['muddat'], 'manba': e['manba'], 'holat': 'faol', 'tartib': e['tartib'],
        'boshlanish': '2026-09-07T09:00:00Z', 'tugash': None, 'rasm_url': None,
    })
js = f"""// Ilm-fan va innovatsiyalar — story-e'lonlar fallback (Supabase ulanmaganda).
// Manba: docs/{SRC_TDTU.name}, docs/{SRC_BIO.name}. GENERATSIYA QILINGAN FAYL — scripts/gen_elonlar.py.

export const INITIAL_ELONLAR = [
{',\n'.join('  ' + json.dumps(r, ensure_ascii=False) for r in js_rows)},
];
"""
OUT_JS.write_text(js, encoding='utf-8')

if __name__ == '__main__':
    print('elonlar:', len(elonlar))
    for e in elonlar:
        print(f" - [{e['turi']}/{e['rang']}] {e['sarlavha'][:70]} | {e['muassasa'][:20]} | {e['qiymat']} | {e['bosqich']} | {e['muddat']}")
