#!/usr/bin/env python3
"""TDTU talabalar ro'yxati (docs/bakalavriat/*.xlsx) + magistr fayli -> Supabase seed + JS fallback.

Ishga tushirish (uttp-next ichidan):  python3 scripts/gen_talaba.py   (openpyxl kerak)

Manbalar:
  docs/bakalavriat/TDTU_03_09_2026_..._Talabalar_ro'yxati.xlsx  - 27 739 talaba (Bakalavr/Magistr/Ordinatura)
  docs/magistr_TDTU_2_3_kurs.xlsx                                - 1 231 magistr (ro'yxatning qism-to'plami):
        mutaxassislik shifri (70910xxx), guruh (kampus/til/qabul yili), 2026/27 kurs raqami.
        Ro'yxatdagi magistr kursi bittaga kam yozilgan (fayl 2-kurs = ro'yxat 1-kurs) — fayl ustun.

Chiqish:
  uttp-next/supabase/seed_talaba.sql   - talaba jadvali + auth hisoblar (login/parol) + profillar
  uttp-next/lib/talaba-fallback.js     - agregat statistikalar (Supabase bo'lmaganda)

Login:  <familiya>.<ism>@talaba.tdtu.uz   (takror bo'lsa raqam qo'shiladi: 2, 3 ...)
Parol:  namuna hisoblar (DEPLOY.md jadvali, 4 nafar) — <ism>2026 (ochiq, login demosi uchun);
        qolgan har bir hisob — base32(sha256('ettp:<login>:<TUZ>'))[:10], kichik harf,
        TUZ = ETTP_PAROL_TUZ muhit o'zgaruvchisi (standart 'ettp-demo-2026' faqat lokal demo;
        bulutga reseed qilishdan oldin MAXFIY qiymat o'rnating).
        Parollar faqat seed SQL'ga yoziladi (crypt bilan) — JS fallbackka namunadan boshqasi tushmaydi.
"""
import base64
import hashlib
import json
import os
import re
import unicodedata
import uuid
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

# Repo ildizi: uttp-next/scripts/gen_talaba.py → ../..
ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs/bakalavriat/TDTU_03_09_2026_holatiga_tahsil_olayotgan_Talabalar_ro'yxati.xlsx"
SRC_MAG = ROOT / 'docs/magistr_TDTU_2_3_kurs.xlsx'
OUT_SQL = ROOT / 'uttp-next/supabase/seed_talaba.sql'
OUT_JS = ROOT / 'uttp-next/lib/talaba-fallback.js'
NAMUNA_JS = ROOT / 'uttp-next/lib/demo-namunalar.js'   # qo'lda yoziladigan kichik modul (sana sinxron)

MANBA_SANASI = '2026-09-03'   # TDTU kontingent ro'yxati holati (docs/bakalavriat/TDTU_03_09_2026_...)

MUASSASA = 'Toshkent davlat tibbiyot universiteti'

# Parol tuzi: bulutga reseed qilishdan oldin ETTP_PAROL_TUZ ga MAXFIY qiymat bering
# (standart qiymat faqat lokal demo uchun). Tuz hech qayerga yozilmaydi.
TUZ = os.environ.get('ETTP_PAROL_TUZ') or 'ettp-demo-2026'


def yashirin_parol(login: str) -> str:
    """Hisobga xos, taxmin qilib bo'lmaydigan parol: base32(sha256('ettp:<login>:<TUZ>'))[:10], kichik harf (a-z, 2-7)."""
    xesh = hashlib.sha256(f'ettp:{login}:{TUZ}'.encode('utf-8')).digest()
    return base64.b32encode(xesh)[:10].decode('ascii').lower()


BOSQICH = {'Bakalavr': 'bakalavr', 'Magistr': 'magistr', 'Ordinatura': 'rezidentura'}


def ascii_slug(s: str) -> str:
    """Faqat a-z: diakritikalar olib tashlanadi, apostrof/tire/boshqa belgilar tushib qoladi."""
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r'[^a-z]', '', s.lower())


def title_case(fish: str) -> str:
    return ' '.join(w[:1].upper() + w[1:].lower() for w in fish.split())


def norm(s) -> str:
    return re.sub(r'\s+', ' ', str(s or '')).strip()


def q(v) -> str:
    """SQL literal."""
    if v is None:
        return 'NULL'
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def guruh_tahlili(g: str):
    """Guruh nomidan kampus / til / qabul yilini ajratadi (topilmasa None)."""
    g = g or ''
    kampus = '2-kampus' if re.search(r'kampus', g, re.I) else ('Bosh bino' if re.search(r'bosh\s*bino', g, re.I) else None)
    til = 'ru' if re.search(r'rus', g, re.I) else ('uz' if re.search(r'\buzb?\b', g, re.I) else None)
    yillar = [int(y) for y in re.findall(r'(20[2-3]\d)', g)]
    if not yillar:
        m = re.search(r'\((2\d)-(2\d)\)', g)  # (24-25)
        if m:
            yillar = [2000 + int(m.group(1))]
    qabul = min(yillar) if yillar else None
    return kampus, til, qabul


# ---------------- Ro'yxat ----------------
ws = openpyxl.load_workbook(SRC, data_only=True, read_only=True).worksheets[0]
rows = [r for r in list(ws.iter_rows(values_only=True))[1:] if r and r[1]]

talabalar = []
login_counter = Counter()
for r in rows:
    tartib, fish, dob, jinsi, fakultet, mutaxassislik, kurs, guruh, talim_turi = r[:9]
    fish = norm(fish)
    # Bo'sh/juda qisqa so'zlar (".", "&", "K") tashlab yuboriladi — birinchi ikkita ma'noli so'z olinadi
    words = [w for w in (ascii_slug(w) for w in fish.split()) if len(w) >= 2]
    familiya = words[0] if words else f't{int(tartib)}'
    ism = words[1] if len(words) > 1 else ''
    base = f'{familiya}.{ism}' if ism else familiya
    login_counter[base] += 1
    n = login_counter[base]
    login = f'{base}{"" if n == 1 else n}@talaba.tdtu.uz'
    tug = datetime.strptime(str(dob).strip(), '%d.%m.%Y').date().isoformat()
    talabalar.append({
        'tartib': int(tartib),
        'user_id': str(uuid.uuid5(uuid.NAMESPACE_URL, f'ettp:talaba:{login}')),
        'fish': title_case(fish),
        'tugilgan_sana': tug,
        'jinsi': str(jinsi).strip().lower(),
        'fakultet': norm(fakultet),
        'mutaxassislik': norm(mutaxassislik),
        'kurs': int(str(kurs).split('-')[0]),
        'guruh': norm(guruh) if guruh else None,
        'talim_turi': str(talim_turi).strip(),
        'login': login,
        'parol': yashirin_parol(login),
        'ism_slug': ism or familiya,   # namuna hisoblarning ochiq paroli uchun
        # magistr faylidan boyitiladi
        'mutaxassislik_kodi': None,
        'kampus': None,
        'til': None,
        'qabul_yili': None,
    })

# ---------------- Magistr fayli bilan boyitish ----------------
ws2 = openpyxl.load_workbook(SRC_MAG, data_only=True, read_only=True).worksheets[0]
mag_rows = [r for r in list(ws2.iter_rows(values_only=True))[1:] if r and r[1]]
kalit = {(t['fish'].upper(), t['tugilgan_sana']): t for t in talabalar if t['talim_turi'] == 'Magistr'}
boyitildi = 0
for r in mag_rows:
    _, fish, dob, jins, guruh, kurs, _turi, kodi, _shakl = r[:9]
    k = (title_case(norm(fish)).upper(), str(dob)[:10])
    t = kalit.get(k)
    if not t:
        raise SystemExit(f'magistr faylidagi talaba ro‘yxatda topilmadi: {fish} {dob}')
    kampus, til, qabul = guruh_tahlili(str(guruh or ''))
    t['kurs'] = int(str(kurs).split('-')[0])          # 2026/27 kursi — fayl ustun
    t['mutaxassislik_kodi'] = str(kodi).strip()
    t['kampus'], t['til'], t['qabul_yili'] = kampus, til, qabul
    boyitildi += 1

assert len({t['login'] for t in talabalar}) == len(talabalar), 'login takror'
assert len({t['user_id'] for t in talabalar}) == len(talabalar), 'uuid takror'
assert all(t['jinsi'] in ('erkak', 'ayol') for t in talabalar)
assert all(t['talim_turi'] in BOSQICH for t in talabalar)
assert all(len(t['parol']) >= 6 for t in talabalar), 'parol juda qisqa'

# Login demosi uchun namuna hisoblar (1-kurs bakalavr, bitiruvchi kurs, magistr, ordinatura).
# Faqat shular ochiq <ism>2026 parolini oladi (DEPLOY.md jadvali va login sahifasi shunga tayanadi).
namuna = [
    next(t for t in talabalar if t['talim_turi'] == 'Bakalavr' and t['kurs'] == 1),
    next(t for t in talabalar if t['talim_turi'] == 'Bakalavr' and t['kurs'] == 6),
    next(t for t in talabalar if t['talim_turi'] == 'Magistr' and t['mutaxassislik_kodi'] and t['kampus']),
    next(t for t in talabalar if t['talim_turi'] == 'Ordinatura'),
]
for n in namuna:
    n['parol'] = f"{n['ism_slug']}2026"

# Shifr → eng ko'p uchragan nom (klassifikator uchun; seed.sql'dagi ro'yxat shu bilan sinxron)
kod_nomlari = defaultdict(Counter)
for t in talabalar:
    if t['mutaxassislik_kodi']:
        kod_nomlari[t['mutaxassislik_kodi']][t['mutaxassislik']] += 1
KOD_NOMI = {k: c.most_common(1)[0][0] for k, c in kod_nomlari.items()}

# ---------------- SQL ----------------
# Bitta buyruq: Supabase CLI seed runner har buyruqni oldindan prepare qiladi, shuning
# uchun TEMP TABLE ishlamaydi. VALUES + data-modifying CTE'lar bilan hammasi bir tranzaksiyada.
COLS = ('tartib', 'user_id', 'fish', 'tugilgan_sana', 'jinsi', 'fakultet', 'mutaxassislik', 'kurs',
        'guruh', 'talim_turi', 'login', 'parol', 'mutaxassislik_kodi', 'kampus', 'til', 'qabul_yili')
vals = ['(' + ', '.join(q(t[k]) for k in COLS) + ')' for t in talabalar]
lines = [
    '-- =========================================================',
    '-- ETTP — TDTU talabalar kontingenti (03.09.2026 holatiga) + auth hisoblar',
    f"-- Manba: docs/bakalavriat/{SRC.name}",
    f"--        docs/{SRC_MAG.name} (magistr boyitish: shifr, kampus, til, qabul yili, kurs — {boyitildi} nafar)",
    f'-- {len(talabalar)} talaba: ' + ', '.join(f'{k} {v}' for k, v in Counter(t['talim_turi'] for t in talabalar).most_common()),
    '-- Login: <familiya>.<ism>@talaba.tdtu.uz. Parol: namuna hisoblar (DEPLOY.md) — <ism>2026;',
    '-- qolganlari — ETTP_PAROL_TUZ dan hosil qilingan 10 belgili yashirin parol (scripts/gen_talaba.py).',
    '-- GENERATSIYA QILINGAN FAYL — scripts/gen_talaba.py; qo‘lda tahrirlanmaydi.',
    '-- Bitta buyruq (CTE): auth.users → auth.identities → public.talaba → public.profiles.',
    '-- Mavjud hisoblar o‘tkazib yuboriladi (idempotent). Magistr shifrlari specializations (seed.sql) bilan mos.',
    '-- =========================================================',
    '',
    f'WITH s({", ".join(COLS)}) AS (',
    '  VALUES',
    ',\n'.join('  ' + v for v in vals),
    '),',
    'yangi AS (',
    '  SELECT s.* FROM s WHERE NOT EXISTS (SELECT 1 FROM auth.users u WHERE u.id = s.user_id::uuid)',
    '),',
    '-- 1. Auth hisoblar (parol bcrypt bilan)',
    'u AS (',
    '  INSERT INTO auth.users (',
    '      instance_id, id, aud, role, email, encrypted_password, email_confirmed_at,',
    '      raw_app_meta_data, raw_user_meta_data, created_at, updated_at,',
    '      confirmation_token, recovery_token, email_change,',
    '      email_change_token_new, email_change_token_current, is_super_admin',
    '  )',
    "  SELECT '00000000-0000-0000-0000-000000000000', y.user_id::uuid, 'authenticated', 'authenticated',",
    "         y.login, extensions.crypt(y.parol, extensions.gen_salt('bf')), NOW(),",
    """         '{"provider":"email","providers":["email"]}', '{}', NOW(), NOW(), '', '', '', '', '', false""",
    '  FROM yangi y',
    '  RETURNING id',
    '),',
    'i AS (',
    '  INSERT INTO auth.identities (id, provider_id, user_id, identity_data, provider, last_sign_in_at, created_at, updated_at)',
    '  SELECT y.user_id::uuid, y.user_id, y.user_id::uuid,',
    "         jsonb_build_object('sub', y.user_id, 'email', y.login, 'email_verified', true, 'phone_verified', false),",
    "         'email', NOW(), NOW(), NOW()",
    '  FROM yangi y',
    '  RETURNING id',
    '),',
    '-- 2. Talaba jadvali',
    't AS (',
    '  INSERT INTO public.talaba (tartib, user_id, fish, tugilgan_sana, jinsi, fakultet, mutaxassislik, kurs, guruh, talim_turi, login,',
    '                             mutaxassislik_kodi, kampus, til, qabul_yili)',
    '  SELECT y.tartib, y.user_id::uuid, y.fish, y.tugilgan_sana::date, y.jinsi, y.fakultet, y.mutaxassislik, y.kurs, y.guruh, y.talim_turi, y.login,',
    '         y.mutaxassislik_kodi, y.kampus, y.til, y.qabul_yili',
    '  FROM yangi y',
    '  RETURNING id',
    ')',
    "-- 3. Profil (rol = 'talaba'): kabinetga kirish va rol aniqlash uchun; magistr shifri klassifikatorga bog'lanadi",
    'INSERT INTO public.profiles (id, fish, tug_ilgan_sana, jinsi, email, hozirgi_bosqich, hozirgi_muassasa, hozirgi_kurs, yonalish_kodi, lavozimi, rol)',
    'SELECT y.user_id::uuid, y.fish, y.tugilgan_sana::date, y.jinsi, y.login,',
    "       CASE y.talim_turi WHEN 'Bakalavr' THEN 'bakalavr' WHEN 'Magistr' THEN 'magistr' ELSE 'rezidentura' END,",
    f"       {q(MUASSASA)}, y.kurs,",
    '       (SELECT sp.kodi FROM public.specializations sp WHERE sp.kodi = y.mutaxassislik_kodi),',
    "       y.talim_turi || ' · ' || y.fakultet, 'talaba'",
    'FROM yangi y;',
    '',
]
OUT_SQL.write_text('\n'.join(lines), encoding='utf-8')

# ---------------- JS fallback ----------------
fak = defaultdict(int)
mut = defaultdict(int)
mag = defaultdict(int)
magmut = defaultdict(int)
yosh = defaultdict(int)
for t in talabalar:
    fak[(t['talim_turi'], t['fakultet'], t['kurs'], t['jinsi'])] += 1
    mut[(t['talim_turi'], t['fakultet'], t['mutaxassislik'], t['mutaxassislik_kodi'])] += 1
    yosh[(t['talim_turi'], int(t['tugilgan_sana'][:4]), t['jinsi'])] += 1
    if t['talim_turi'] == 'Magistr':
        mag[(t['kurs'], t['jinsi'], t['kampus'], t['til'], t['qabul_yili'])] += 1
        magmut[(t['mutaxassislik_kodi'], t['mutaxassislik'], t['kurs'], t['jinsi'])] += 1

fak_rows = [
    {'talim_turi': k[0], 'fakultet': k[1], 'kurs': k[2], 'jinsi': k[3], 'soni': v}
    for k, v in sorted(fak.items())
]
mut_rows = [
    {'talim_turi': k[0], 'fakultet': k[1], 'mutaxassislik': k[2], 'mutaxassislik_kodi': k[3], 'soni': v}
    for k, v in sorted(mut.items(), key=lambda kv: (kv[0][0], kv[0][1], -kv[1]))
]
mag_rows_js = [
    {'kurs': k[0], 'jinsi': k[1], 'kampus': k[2], 'til': k[3], 'qabul_yili': k[4], 'soni': v}
    for k, v in sorted(mag.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2] or '', kv[0][3] or '', kv[0][4] or 0))
]
yosh_rows = [
    {'talim_turi': k[0], 'tugilgan_yil': k[1], 'jinsi': k[2], 'soni': v}
    for k, v in sorted(yosh.items())
]
magmut_rows = [
    {'mutaxassislik_kodi': k[0], 'mutaxassislik': k[1], 'kurs': k[2], 'jinsi': k[3], 'soni': v}
    for k, v in sorted(magmut.items(), key=lambda kv: (kv[0][0] or '', kv[0][1], kv[0][2], kv[0][3]))
]

def js_rows(rows):
    return ',\n'.join('  ' + json.dumps(r, ensure_ascii=False) for r in rows)


js = f"""// TDTU talabalar kontingenti — agregat fallback (Supabase ulanmaganda).
// Manba: docs/bakalavriat/{SRC.name} + docs/{SRC_MAG.name}
// {len(talabalar)} talaba. GENERATSIYA QILINGAN FAYL — scripts/gen_talaba.py.
// Ustunlar supabase/migrations/006_talaba.sql va 007_magistr.sql dagi view'larga mos.

export const TALABA_MANBA_SANASI = '{MANBA_SANASI}';

// talaba_fakultet_stat: talim_turi × fakultet × kurs × jinsi
export const INITIAL_TALABA_FAKULTET_STAT = [
{js_rows(fak_rows)},
];

// talaba_mutaxassislik_stat: talim_turi × fakultet × mutaxassislik × shifr (shifr faqat magistrda)
export const INITIAL_TALABA_MUTAXASSISLIK_STAT = [
{js_rows(mut_rows)},
];

// talaba_magistr_stat: magistr — kurs × jinsi × kampus × til × qabul yili
export const INITIAL_TALABA_MAGISTR_STAT = [
{js_rows(mag_rows_js)},
];

// talaba_yosh_stat: talim_turi × tug'ilgan yil × jinsi
export const INITIAL_TALABA_YOSH_STAT = [
{js_rows(yosh_rows)},
];

// talaba_magistr_mut_stat: magistr — shifr × mutaxassislik × kurs × jinsi
export const INITIAL_TALABA_MAGISTR_MUT_STAT = [
{js_rows(magmut_rows)},
];

// Login demosi uchun namuna hisoblar (seed_talaba.sql bilan sinxron). Faqat shu 4 hisobning
// ochiq <ism>2026 paroli shu yerda; qolgan parollar hech qachon JS'ga yozilmaydi.
export const DEMO_TALABALAR = [
{js_rows([{'email': n['login'], 'parol': n['parol'], 'ism': n['fish'], 'bosqich': f"{n['talim_turi']} {n['kurs']}-kurs · {n['fakultet']}", 'bosqichId': BOSQICH[n['talim_turi']], 'kurs': n['kurs']} for n in namuna])},
];
"""
# JS'ga faqat namuna parollari tushganini tekshirish (yashirin parollar sizmasin)
assert js.count('"parol"') == len(namuna), 'JS fallbackda ortiqcha parol'
assert all(n['parol'].endswith('2026') for n in namuna)
OUT_JS.write_text(js, encoding='utf-8')

# lib/demo-namunalar.js (TalabaProfil, dashboard) shu sanani o'z konstantasida takrorlaydi — u katta fallback
# modulni import qilmasligi uchun. Sana o'zgarsa u faylni ham qo'lda yangilash shart.
assert f"TALABA_MANBA_SANASI = '{MANBA_SANASI}'" in NAMUNA_JS.read_text(encoding='utf-8'), (
    f'lib/demo-namunalar.js dagi TALABA_MANBA_SANASI {MANBA_SANASI} emas — qo‘lda yangilang'
)

if __name__ == '__main__':
    print('parol tuzi:', 'ETTP_PAROL_TUZ (maxfiy)' if os.environ.get('ETTP_PAROL_TUZ') else 'standart (faqat lokal demo!)')
    print('talabalar:', len(talabalar), '| magistr boyitildi:', boyitildi, '| shifrlar:', len(KOD_NOMI))
    print('sql bytes:', OUT_SQL.stat().st_size, '| js bytes:', OUT_JS.stat().st_size)
    print('fak rows:', len(fak_rows), '| mut rows:', len(mut_rows), '| mag rows:', len(mag_rows_js), '| magmut rows:', len(magmut_rows), '| yosh rows:', len(yosh_rows))
    print('magistr kurs:', Counter(t['kurs'] for t in talabalar if t['talim_turi'] == 'Magistr'))
    print('kampus:', Counter(t['kampus'] for t in talabalar if t['talim_turi'] == 'Magistr'))
    print('til:', Counter(t['til'] for t in talabalar if t['talim_turi'] == 'Magistr'))
    print('qabul:', Counter(t['qabul_yili'] for t in talabalar if t['talim_turi'] == 'Magistr'))
    print('namuna:', [(n['login'], n['parol']) for n in namuna])
